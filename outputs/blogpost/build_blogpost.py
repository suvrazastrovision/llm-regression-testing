from pathlib import Path
import ast
import hashlib
import json
import shutil
import zipfile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
BUNDLE = OUT / 'walkthrough'
BUNDLE.mkdir(exist_ok=True)
for name in ['drug_discovery.py', 'requirements.txt', 'pyproject.toml', 'uv.lock',
             '.env.example', '.gitignore', 'LICENSE', 'README.md',
             'docs/README.md', 'examples/sample_results.json', 'examples/README.md',
             'tests/test_drug_discovery.py']:
    target = BUNDLE / name
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ROOT / name, target)

INSPECTOR = '''"""Inspect a complete result file without model requests or trace uploads."""
import argparse
import json
from pathlib import Path

from drug_discovery import print_report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    args = parser.parse_args()
    rows = json.loads(args.file.read_text(encoding="utf-8-sig"))
    if not isinstance(rows, list) or not rows:
        parser.error("Expected a nonempty list of result records.")
    trials = {row.get("trial", 1) for row in rows}
    if any(type(trial) is not int or trial < 1 for trial in trials):
        parser.error("Trial identifiers must be positive integers.")
    modes = {row.get("mode", "demo") for row in rows}
    if len(modes) != 1 or not modes <= {"demo", "comparison"}:
        parser.error("Expected one valid mode per file.")
    try:
        gate = print_report(rows, repeats=max(trials), mode=modes.pop())
    except (ValueError, TypeError) as error:
        parser.error(f"Incomplete or invalid scored run: {error}")
    for row in rows:
        print(f"\\n{row['version']} trial {row.get('trial', 1)}")
        print("Answer:", row["answer"])
        print("Judge reason:", row["reason"])
    return 0 if gate["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
'''
(BUNDLE / 'inspect_results.py').write_text(INSPECTOR, encoding='utf-8')
SMOKE = '''"""Small deterministic checks of gate behavior; no API calls."""
import contextlib
import io
import unittest

from drug_discovery import TESTS, print_report


class GateChecks(unittest.TestCase):
    def gate(self, accuracy_a, accuracy_b):
        rows = [
            dict(question=q, version=v, trial=1,
                 accuracy=score, clarity=2, reason="Fixture")
            for q, _ in TESTS
            for v, score in (("A", accuracy_a), ("B", accuracy_b))
        ]
        with contextlib.redirect_stdout(io.StringIO()):
            return print_report(rows)["passed"]

    def test_complete_answers_pass(self):
        self.assertTrue(self.gate(2, 2))

    def test_accuracy_drop_fails(self):
        self.assertFalse(self.gate(2, 1))

    def test_equal_major_errors_fail(self):
        self.assertFalse(self.gate(0, 0))

    def test_shared_incompleteness_passes_relative_gate(self):
        self.assertTrue(self.gate(1, 1))


if __name__ == "__main__":
    unittest.main()
'''
(BUNDLE / 'check_gate.py').write_text(SMOKE, encoding='utf-8')
manifest = {
    'source_commit': '00b0f0315cab6047f1e2f366e7cc57930cf479a3',
    'prepared_on': '2026-10-09',
    'main_script_sha256': hashlib.sha256((BUNDLE / 'drug_discovery.py').read_bytes()).hexdigest(),
    'validation': '20 existing offline tests passed; live API calls not run for this revision',
}
(BUNDLE / 'walkthrough_manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
(BUNDLE / 'WALKTHROUGH.txt').write_text('''LLM regression testing walkthrough
Read the accompanying revised Word article for the scientific rationale.
From this extracted folder in PowerShell:
python -m venv .venv
.\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt
.\\.venv\\Scripts\\python.exe -m pip check
.\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v
.\\.venv\\Scripts\\python.exe check_gate.py
.\\.venv\\Scripts\\python.exe inspect_results.py examples/sample_results.json
The sample inspector exits 2 because the published sample fails its gate.
For locked installation instead: uv sync --locked
For live runs: copy .env.example to .env only if .env is absent, fill your
OpenAI and Langfuse credentials, then run:
.\\.venv\\Scripts\\python.exe drug_discovery.py --mode demo
.\\.venv\\Scripts\\python.exe drug_discovery.py --mode comparison --repeats 3
Live requests incur API usage and upload prompts/answers to Langfuse.
Exit codes: 0 passing gate; 2 failed gate (or invalid CLI); 1 runtime error.
No actual .env, local results, or virtual environment is included.
''', encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

# compact_reference_guide; named overrides: Code, Reference, Equation,
# EditorialTitle (a compact editorial opening without a separate cover).
for name, size, color, before, after in [
    ('Normal', 11, '202020', 0, 6),
    ('Title', 25, '17365D', 0, 8),
    ('Subtitle', 12, '555555', 0, 8),
    ('Heading 1', 16, '2E74B5', 18, 10),
    ('Heading 2', 13, '2E74B5', 14, 7),
    ('Heading 3', 12, '1F4D78', 10, 5),
]:
    st = doc.styles[name]
    st.font.name, st.font.size, st.font.color.rgb = 'Calibri', Pt(size), RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.line_spacing = 1.25
    if name.startswith('Heading'):
        st.font.bold = True
        st.paragraph_format.keep_with_next = True
for name, font, size, spacing, after in [
    ('Code', 'Consolas', 8.5, 10.5, 0),
    ('Reference', 'Calibri', 10, 12, 6),
    ('Equation', 'Consolas', 10, 13, 6),
]:
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name, st.font.size = font, Pt(size)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.line_spacing = Pt(spacing)
    st.paragraph_format.keep_together = True
    if name == 'Code':
        st.paragraph_format.widow_control = False
        st.font.color.rgb = RGBColor.from_string('202020')
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F3F5F7')
        st.element.get_or_add_pPr().append(shade)
header = sec.header.paragraphs[0]
header.text = 'LLM regression testing | Scientific walkthrough'
header.style = doc.styles['Reference']
header.runs[0].font.color.rgb = RGBColor.from_string('666666')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Page ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for r in footer.runs:
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string('666666')

def p(text, style=None):
    return doc.add_paragraph(text, style=style)

def h(text, level=1):
    doc.add_heading(text, level)

def code(text):
    # One logical source line per Word paragraph. Word may visually wrap a
    # line, but the source text and indentation remain intact.
    for line in text.rstrip('\n').split('\n'):
        p(line, 'Code')
    p('')

def link(paragraph, label, url):
    part = paragraph.part
    rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
    el = OxmlElement('w:hyperlink'); el.set(qn('r:id'), rid)
    run = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); prop.append(color)
    under = OxmlElement('w:u'); under.set(qn('w:val'), 'single'); prop.append(under)
    run.append(prop); t = OxmlElement('w:t'); t.text = label; run.append(t)
    el.append(run); paragraph._p.append(el)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    pr = t._tbl.tblPr
    tw = pr.find(qn('w:tblW')); tw.set(qn('w:w'), '9360'); tw.set(qn('w:type'), 'dxa')
    ti = OxmlElement('w:tblInd'); ti.set(qn('w:w'), '120'); ti.set(qn('w:type'), 'dxa'); pr.append(ti)
    margins = OxmlElement('w:tblCellMar')
    for name, value in [('top',80),('bottom',80),('start',120),('end',120)]:
        el = OxmlElement('w:'+name); el.set(qn('w:w'),str(value)); el.set(qn('w:type'),'dxa'); margins.append(el)
    pr.append(margins)
    borders = OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:'+name); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'CCD5DF'); borders.append(el)
    pr.append(borders)
    grid = t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'),str(width)); grid.append(col)
    for c, text in zip(t.rows[0].cells, headers): c.text = text
    for values in rows:
        for c, text in zip(t.add_row().cells, values): c.text = str(text)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, row in enumerate(t.rows):
        for c, width in zip(row.cells, widths):
            c.width = Inches(width / 1440)
            w = c._tc.get_or_add_tcPr().find(qn('w:tcW')); w.set(qn('w:w'),str(width)); w.set(qn('w:type'),'dxa')
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1.1
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.bold = index == 0
            if index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF5'); c._tc.get_or_add_tcPr().append(shade)
    p('')

p('Regression Testing for Scientific LLM Answers', 'Title')
p('An evidence-based walkthrough with complete Python code', 'Subtitle')
p('Suvra Nath | Correspondence: suvranath047@gmail.com')

h('Abstract')
p('A language model can produce a fluent explanation while reversing a biological mechanism or extending an animal finding beyond its evidence. This walkthrough implements a small regression-testing workflow: generate answers, score them against literature-derived reference facts, compare prompt versions, and retain an auditable record. Two questions concerning ventral tegmental area (VTA) circuitry illustrate the method. The default experiment is a deliberate-error control; a separate mode compares two plausible prompts. The published sample fails the regression gate but also contains a judge error. The workflow therefore demonstrates software behavior and error detection in specific cases; it does not establish the scientific validity of an automated evaluator or the superiority of a prompt across a research domain.')

h('Start with a testable claim')
p('The operational question is: under a fixed question set, rubric, answer model, and judge configuration, does a candidate prompt produce lower rubric scores than a baseline? Here, regression means a decrease according to a declared software rule. Factual accuracy is operationalized as agreement with the supplied reference facts, including experimental context and limits of interpretation. This is narrower than scientific truth or literature completeness.')
p('An answer is the observational unit. The two questions are the content units; repeated calls are repeated measurements within those units. Fresh API conversations avoid shared chat history but do not establish statistical independence of model errors. The evaluator is another model request. Its output is a fallible measurement that requires comparison with expert labels. Research on LLM judges documents position, verbosity, and self-enhancement biases; evidence from general chat benchmarks does not validate this neuroscience rubric. [6]')
p('The complete implementation is reproduced in Appendix A. Copyable companion files are provided in LLM_Regression_Testing_Walkthrough.zip. The main script is preserved from the local repository snapshot identified in walkthrough_manifest.json. [1]')

h('Why these brain questions?')
h('KCNQ channels and circuit specificity', 2)
p('Friedman and colleagues examined KCNQ/Kv7 potassium channels in a chronic social defeat stress mouse model. Channel opening reduces excitability through a stabilizing potassium conductance. Retigabine and KCNQ upregulation reduced excessive VTA dopamine-neuron activity and reversed the measured depression-related behaviors in susceptible mice. Projection-targeted KCNQ3 expression produced behavioral effects in the VTA-to-nucleus-accumbens pathway, whereas corresponding expression in the VTA-to-medial-prefrontal-cortex pathway did not. [2]')
p('These results motivate a circuit-dependent target hypothesis. The intervention and behavioral assays in mice do not, by themselves, demonstrate human antidepressant efficacy, a clinical cure, or a universal dopamine mechanism of depression. The reference answer preserves this distinction. A response that calls retigabine a KCNQ blocker reverses the intervention and is a major mechanistic error.')
h('Ketamine: reinforcement and measured plasticity', 2)
p('Simmler and colleagues reported in mice that ketamine antagonized N-methyl-D-aspartate (NMDA) receptors on VTA gamma-aminobutyric acid (GABA) neurons, disinhibiting dopamine neurons. Dopamine D2 receptor feedback curtailed the dopamine response. Ketamine supported reinforcement, while brief dopamine transients together with NMDA-receptor antagonism prevented the forms of synaptic plasticity examined in the VTA and nucleus accumbens. [3]')
p('The study motivates evaluating response duration and plasticity alongside reinforcement. It does not establish an absence of addiction risk in humans or provide a complete account of ketamine antidepressant action. Reinforcement, the measured plasticity endpoints, and human addiction are distinct outcomes. This article makes no treatment recommendation.')

h('Distinguish the two experimental modes')
table(['Mode', 'Prompt A', 'Prompt B'], [
    ['demo (default)', 'Receives the reference facts', 'Asked to assert deliberately false claims'],
    ['comparison', 'Basic scientific-answer prompt', 'Mechanisms, abbreviations, uncertainty, and translation limits'],
], [1800, 3500, 4060])
p('The demo asks whether the pipeline can detect conspicuous controlled errors. Its prompts differ in both information access and correctness instructions, so its scores cannot estimate the benefit of an improved prompt. The expected A-pass/B-fail pattern is a design expectation; the model may correct the false premise, and the judge may misgrade it. Scores are not forced.')
p('Comparison mode withholds reference facts from both answer prompts and gives the same reference facts to the judge. It is a more relevant baseline/candidate comparison, although the structured candidate may be easier for this rubric to reward. Neither mode is a randomized user A/B experiment. No research paper is retrieved by the script; the keys are manually supplied summaries. Both prompts request at most 100 words, which is an instruction rather than an enforced word-count constraint.')

h('Get the project ready')
p('Use Python 3.10 or later within the project range (less than 4), PowerShell, an OpenAI API account, and a Langfuse project. Either extract the companion bundle into a working folder, or clone the repository. Run the commands from that folder. If cloning, use:')
code('git clone https://github.com/suvrazastrovision/llm-regression-testing.git\ncd llm-regression-testing')
p('Create a virtual environment and install the declared requirements:')
code('python -m venv .venv\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt\n.\\.venv\\Scripts\\python.exe -m pip check')
p('The entire requirements.txt is:')
code((ROOT / 'requirements.txt').read_text(encoding='utf-8-sig'))
p('These version ranges are compatible declarations, not exact pins. For an existing project or the companion bundle, uv sync --locked uses the included pyproject.toml and uv.lock to install the locked environment. [7] Do not combine installation methods for the same environment without understanding the resulting dependency changes. A lock file improves environment reproducibility; it does not make hosted model outputs identical.')
code('uv sync --locked')

h('Add your keys once')
p('Create .env only if it does not already exist. Preserve an existing configuration:')
code('if (-not (Test-Path -LiteralPath .env)) {\n    Copy-Item -LiteralPath .env.example -Destination .env\n}')
p('Replace the placeholders in .env with your own credentials. The following complete configuration retains the repository model defaults; model access depends on your account:')
code('OPENAI_API_KEY=your_actual_openai_key\nLLM_MODEL=gpt-4o-mini\nJUDGE_MODEL=gpt-4o-mini\nLANGFUSE_PUBLIC_KEY=your_langfuse_public_key\nLANGFUSE_SECRET_KEY=your_langfuse_secret_key\nLANGFUSE_BASE_URL=https://cloud.langfuse.com')
p('OpenAI and Langfuse use separate credentials. Set the Langfuse base URL for your project region or self-hosted instance; the example is the EU cloud endpoint. [4] The script loads .env beside drug_discovery.py. Existing process environment variables take precedence. Git ignores .env; the bundle includes only .env.example.')

h('Check the plumbing before live requests')
code('.\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v\n.\\.venv\\Scripts\\python.exe check_gate.py')
p('The supplied repository suite mocks model and tracing clients and covers aggregation, exit codes, invalid responses, and partial-answer preservation. All 20 existing tests passed during preparation of this revision. Appendix C gives a smaller, complete set of gate checks for readers building directly from the article. These are software tests; they neither verify the papers nor measure judge agreement with scientists. No live API experiment was performed for this revision.')

h('Run the deliberate-error demo')
code('.\\.venv\\Scripts\\python.exe drug_discovery.py --mode demo\n$LASTEXITCODE')
p('With two questions and one trial, a complete run makes four generation requests and four judge requests. In general, Q questions and R repeats make 2QR requests of each type, or 4QR total calls before SDK retries. Calls incur OpenAI API usage. The script prints the answers and scores, saves timestamped JSON under results/, and queues tracing data for Langfuse.')
p('Open your Langfuse project and inspect traces named drug-discovery-A and drug-discovery-B. The wrapped OpenAI client records generate-answer and score-answer calls; the enclosing span adds run, mode, version, and trial metadata, and the script attaches numeric scores. Langfuse receives prompts, answers, usage, and scores. flush() attempts to send queued data before exit. [4] A tracing-upload failure can leave the local gate result intact, so separately verify remote trace completeness.')

h('Meet the judge and the decision rule')
p('The judge receives the question, reference facts, and candidate answer without an explicit A/B label. This removes a direct label cue but cannot conceal differences in answer style or prevent biased assessment. Pydantic constrains accuracy and clarity to 0, 1, or 2. Structured output enforces the expected response shape; it does not establish the truth of the assigned score. [5] Refusals, absent parsed scores, and empty explanations from the judge invalidate the run.')
table(['Score', 'Accuracy', 'Clarity'], [
    ['0', 'Major error, contradiction, refusal, or irrelevant answer', 'Confusing'],
    ['1', 'Correct but incomplete', 'Understandable; technical terms unexplained'],
    ['2', 'All required facts; no contradictions', 'Clear reasoning; essential abbreviations defined'],
], [700, 4860, 3800])
p('Individual answer PASS requires accuracy = 2. Clarity is reported separately. Let s(q,v,r) denote the accuracy score for question q, version v, trial r. The per-question mean is the sum across trials divided by R. The gate fails if any question has a lower candidate mean, or any candidate trial scores zero:')
p('FAIL if any q: mean(q,B) < mean(q,A)\n     or any q,r: s(q,B,r) = 0', 'Equation')
p('The rule is a policy threshold, not a significance test. Scores are ordinal; taking their mean additionally treats the category spacing as equal for this engineering summary. A mean difference is therefore not a calibrated effect size. Two incomplete versions scoring 1 can pass the relative gate while both fail their individual answer checks. A single candidate score of 0 vetoes the run regardless of its mean.')
p('Exit 0 means the completed gate passed; exit 2 means it failed; exit 1 indicates a handled configuration, request, validation, or file error. Invalid command-line arguments also exit 2. A failing demo exits 2 even when error detection is expected. Generation refusals are retained as answers for accuracy grading; empty or length-truncated generation responses are rejected. Judge truncation and content filtering invalidate the comparison.')

h('Read the public sample honestly')
p('The published demo sample was recorded on 8 October 2026 at 16:28:12 UTC, using gpt-4o-mini for generation and judging with one trial. The following are its original model-assigned scores, not new measurements or expected outputs for every run. [1]')
table(['Question', 'Version', 'Accuracy', 'Clarity'], [
    ['KCNQ / VTA', 'A', '2', '2'], ['KCNQ / VTA', 'B', '1', '2'],
    ['Ketamine / VTA', 'A', '2', '2'], ['Ketamine / VTA', 'B', '0', '1'],
], [3900, 1500, 1980, 1980])
p('Across the two questions, original mean accuracy is 2 for A and 0.5 for B, giving B minus A = -1.5; the gate fails. The KCNQ B answer reverses retigabine action and firing direction. Under the declared rubric, that contradiction merits accuracy 0 rather than the assigned 1. The repository review notes disclose this inconsistency while preserving the raw judgment. The gate would still fail after that correction, but its numerical summaries would change.')
p('Preserve raw model scores and store reviewer adjudications as a separate record with the rationale and reviewer provenance. This one disagreement demonstrates that the judge can misapply the rubric; four outputs cannot estimate its general error rate or establish prompt superiority.')
p('Appendix B provides a complete offline inspector. To inspect the supplied sample after extracting the bundle, run:')
code('.\\.venv\\Scripts\\python.exe inspect_results.py examples/sample_results.json')
p('The inspector prints the reconstructed gate, answers, and judge reasons and exits 2 for this sample. For new runs, pass the exact results filename printed by drug_discovery.py. It rejects incomplete scored runs rather than aggregating them. Its imports require the installed dependencies, but it neither requests answers nor uploads traces.')

h('Compare two useful prompts')
code('.\\.venv\\Scripts\\python.exe drug_discovery.py --mode comparison --repeats 3')
p('This run makes twelve answer requests and twelve judge requests before retries. A/B order alternates across questions and trials. Each repeated answer is judged once, so observed variation combines generation and evaluation noise; the code does not independently rejudge a fixed answer. Keep the judge and rubric fixed when examining a prompt change. Repeated calls to two questions increase observations within these questions, not coverage of the scientific domain.')
p('Record the code revision, dependency lock, Python version, execution time, model identifiers, complete prompts, reference facts, and rubric. The script saves the prompts, facts, identifiers, mode, trials, and reasons. It does not explicitly persist every provider response identifier, model snapshot, usage field, or sampling setting in local JSON. Archive those separately if required for an expanded experiment. A model alias alone does not establish immutable model behavior.')

h('From demonstration to scientific evaluation')
p('A defensible extension starts with a predefined, independently sourced question set spanning the intended use. Specify inclusion rules, required facts, contradiction criteria, handling of incomplete answers, and an acceptance threshold before assessing the candidate. Freeze a development set for prompt design and a held-out evaluation set to reduce tuning to the test items. The present two cases are a tutorial selection and provide no population-level estimate.')
p('Have at least two qualified reviewers score representative correct, incomplete, and incorrect answers independently, blinded to prompt version. Adjudicate disagreements and retain both initial labels and final rationale. Compare judge labels with those human labels using a confusion matrix, exact agreement, and detection of major errors. For ordinal scores, weighted agreement statistics may be useful if the weighting rule is declared. Report uncertainty and disagreements, rather than treating human adjudication as error-free.')
p('In a larger comparison, report question-level outcomes, score distributions, complete-answer proportions, and the frequency of major errors before any aggregate. If estimating uncertainty, account for repeated observations nested within questions; a question-level paired analysis or cluster bootstrap needs a sufficiently broad question sample. Three calls per question do not justify a claim of statistical significance. To characterize evaluation noise separately, rejudge fixed answers under a predefined protocol.')
p('Keep incomplete runs separate from completed comparisons. The script attempts to save collected answers even when later judging fails, but filesystem failure can still prevent saving. Inspect the error and local file before resuming. Offline tests do not certify factual reliability, and a relative gate PASS does not establish fitness for scientific or clinical use.')

h('Interpretation')
p('This workflow makes prompt changes inspectable through predefined reference facts, structured scores, a declared gate, and retained outputs. The reviewed sample shows a candidate failure and a rubric-application error by the evaluator. Its supported conclusion is limited to these observed cases and the tested software behavior. Broader claims require domain coverage, independent reviewer calibration, and an analysis plan that separates generation variation from measurement error.')

h('References')
refs = [
    ('[1] Nath, S. LLM Regression Testing. Source repository and reviewed public sample. Local source commit 00b0f0315cab6047f1e2f366e7cc57930cf479a3.', 'https://github.com/suvrazastrovision/llm-regression-testing'),
    ('[2] Friedman, A. K., Juarez, B., Ku, S. M., et al. (2016). KCNQ channel openers reverse depressive symptoms via an active resilience mechanism. Nature Communications, 7, 11671. doi:10.1038/ncomms11671.', 'https://www.nature.com/articles/ncomms11671'),
    ('[3] Simmler, L. D., Li, Y., Hadjas, L. C., Hiver, A., van Zessen, R., & Luscher, C. (2022). Dual action of ketamine confines addiction liability. Nature, 608, 368-373. doi:10.1038/s41586-022-04993-7.', 'https://www.nature.com/articles/s41586-022-04993-7'),
    ('[4] Langfuse. Observability for OpenAI SDK (Python). Official integration documentation.', 'https://langfuse.com/integrations/model-providers/openai-py'),
    ('[5] OpenAI. Structured model outputs. Official OpenAI API documentation.', 'https://developers.openai.com/api/docs/guides/structured-outputs'),
    ('[6] Zheng, L., Chiang, W.-L., Sheng, Y., et al. (2023). Judging LLM-as-a-Judge with MT-Bench and Chatbot Arena. arXiv:2306.05685.', 'https://arxiv.org/abs/2306.05685'),
    ('[7] Astral. Locking and syncing. Official uv documentation.', 'https://docs.astral.sh/uv/concepts/projects/sync/'),
]
for text, url in refs:
    para = p(text + ' ', 'Reference'); link(para, 'Source', url)
p('Software documentation checked on 9 October 2026. Scientific references define the two selected case studies; this is not a systematic literature review.', 'Reference')

doc.add_page_break()
h('Appendix A. Complete main script')
p('Save the following full listing as drug_discovery.py. It is the repository implementation used for this walkthrough, with no omitted functions. Both deliberate false claims appear solely as evaluation controls. The companion bundle provides the exact source file so that execution does not depend on copying code from Word.')
source = (ROOT / 'drug_discovery.py').read_text(encoding='utf-8-sig')
code(source)
doc.add_page_break()
h('Appendix B. Complete result inspector')
p('Save as inspect_results.py beside drug_discovery.py. The known questions and gate come from that main script. Do not combine records from different modes or runs in one file.')
code(INSPECTOR)
h('Appendix C. Complete offline gate checks')
p('Save as check_gate.py beside drug_discovery.py. These fixtures test the distinction between a relative comparison and complete scientific answers. The full twenty-test repository suite is also included in the companion bundle.')
code(SMOKE)

doc.core_properties.title = 'Regression Testing for Scientific LLM Answers'
doc.core_properties.subject = 'Scientific walkthrough with complete executable code'
doc.core_properties.author = 'Suvra Nath'
doc.core_properties.keywords = 'LLM evaluation, regression testing, VTA, Langfuse, OpenAI'
path = OUT / 'LLM_Regression_Testing_Revised.docx'
doc.save(path)

# Check exact code round trips: no shortened listings or typography changes.
loaded = Document(path)
actual = '\n'.join(para.text for para in loaded.paragraphs if para.style.name == 'Code')
for name, text in [('main', source), ('inspector', INSPECTOR), ('gate checks', SMOKE)]:
    assert text.rstrip('\n') in actual, f'{name} code did not round-trip'
    ast.parse(text)
assert len(loaded.tables) == 3
assert all(sum(int(col.get(qn('w:w'))) for col in t._tbl.tblGrid) == 9360 for t in loaded.tables)
assert sec.page_width.inches == 8.5 and sec.left_margin.inches == 1

zip_path = OUT / 'LLM_Regression_Testing_Walkthrough.zip'
with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as archive:
    for item in sorted(BUNDLE.rglob('*')):
        if item.is_file():
            archive.write(item, item.relative_to(BUNDLE).as_posix())
print(f'Created {path.name} and {zip_path.name}')
print(f'Full main listing: {len(source.splitlines())} lines; exact code round-trip passed')
print('Preset: compact_reference_guide; header: compact editorial opening')
